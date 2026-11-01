from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

path = r"C:\Users\DELL\Dropbox\PC\Downloads\HQT\permission_matrix.xlsx"
wb = load_workbook(path)
if "Phân loại SQL" in wb.sheetnames:
    del wb["Phân loại SQL"]
ws = wb.create_sheet("Phân loại SQL")
rows = [
    ["File", "Tác dụng", "Module"],
    ["01_database_init.sql", "Script bài tập gốc: tạo lại database, bảng, dữ liệu mẫu, view, procedure và function; có lệnh xóa database nên không chạy trên dữ liệu đang dùng.", "1 + 2"],
    ["02_schema_upgrade.sql", "Nâng cấp an toàn PhoneStore, bổ sung cột, view, function, procedure và cấu trúc cần cho web; không xóa đơn hàng hiện có.", "1"],
    ["03_catalog_seed.sql", "Thêm dữ liệu danh mục, thương hiệu và sản phẩm nền ban đầu.", "2"],
    ["09_report_queries.sql", "Các truy vấn báo cáo: sản phẩm đang bán, tồn kho thấp, doanh thu, đơn hàng và khách hàng.", "3"],
    ["91_cleanup_legacy.sql", "Dọn các SKU cũ đã xác định; nếu đã phát sinh giao dịch thì chuyển sang ngừng bán thay vì xóa.", "2"],
    ["06_product_filters.sql", "Tạo procedure usp_FilterProducts để lọc sản phẩm theo danh mục, thương hiệu và từ khóa.", "3"],
    ["04_catalog_expand.sql", "Bổ sung catalog mở rộng kèm giá, tồn kho, ảnh, bảo hành và nguồn dữ liệu.", "2"],
    ["07_checkout.sql", "Tạo procedure usp_CreateOrder, đặt hàng, trừ tồn kho và ghi Orders/OrderItems trong transaction.", "4 + 5 + 6"],
    ["90_demo_accounts.sql", "Tạo tài khoản mẫu admin và user_demo để đăng nhập kiểm thử.", "2 + 7"],
    ["12_order_sync_check.sql", "Kiểm tra web đặt hàng đã đồng bộ dữ liệu vào SQL hay chưa.", "3 + 5"],
    ["10_security_roles.sql", "Tạo role quản trị viên, nhân viên, người xem báo cáo; cấp quyền bằng GRANT và hạn chế bằng DENY.", "7"],
    ["11_security_tests.sql", "Kiểm thử quyền của từng role, xác nhận thao tác được phép và bị từ chối.", "7"],
]
for row in rows:
    ws.append(row)
header_fill = PatternFill("solid", fgColor="7C8CFF")
header_font = Font(color="FFFFFF", bold=True, size=12)
thin = Side(style="thin", color="B9BFE5")
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(bottom=thin)
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 105
ws.column_dimensions["C"].width = 16
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.row_dimensions[1].height = 28
for i in range(2, ws.max_row + 1):
    ws.row_dimensions[i].height = 44
wb.save(path)
print(path)
