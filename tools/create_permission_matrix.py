from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rows = [
    ["Vai trò / tài khoản", "Đọc sản phẩm", "Sửa sản phẩm", "Xóa sản phẩm", "Đọc đơn hàng", "Tạo đơn hàng", "Xem báo cáo", "Nhập kho"],
    ["Quản trị viên (role_store_admin)", "Cho phép", "Cho phép", "Không cấp trực tiếp", "Cho phép qua view", "Cho phép", "Cho phép", "Cho phép"],
    ["Nhân viên (role_store_staff)", "Cho phép qua view", "Không cấp trực tiếp", "Từ chối", "Cho phép qua view", "Qua thủ tục", "Cho phép qua view", "Qua thủ tục"],
    ["Người xem báo cáo (role_report_viewer)", "Cho phép qua view", "Từ chối", "Từ chối", "Cho phép qua view", "Từ chối", "Cho phép", "Từ chối"],
    ["Ứng dụng web (phone_app)", "Cho phép qua view", "Từ chối", "Từ chối", "Cho phép qua view", "Thủ tục tạo đơn", "Cho phép dashboard", "Từ chối"],
]

wb = Workbook()
ws = wb.active
ws.title = "Permission Matrix"
for row in rows:
    ws.append(row)
header_fill = PatternFill("solid", fgColor="7C8CFF")
header_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="D9DDF4")
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
for col in range(1, ws.max_column + 1):
    ws.column_dimensions[get_column_letter(col)].width = 23 if col > 1 else 24
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions
ws.row_dimensions[1].height = 32
wb.save(r"C:\Users\DELL\Dropbox\PC\Downloads\HQT\permission_matrix.xlsx")
print("created")
